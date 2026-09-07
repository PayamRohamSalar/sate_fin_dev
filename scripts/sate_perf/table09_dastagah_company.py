# table09_dastagah_company.py
from chart_config import (
    load_data, BASE_DIR
)
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
import pandas as pd

TABLES_DIR = BASE_DIR / "tables"
TABLES_DIR.mkdir(parents=True, exist_ok=True)

df = load_data()
df_1405 = df[df["سال"] == 1405].copy()
df_1405 = df_1405.dropna(subset=["دستگاه اجرایی مرتبط", "نام مشمول", "اعتبار"])

# ── آماده‌سازی داده ──────────────────────────────────────────────────────────
by_company = (
    df_1405.groupby(["دستگاه اجرایی مرتبط", "نام مشمول"])["اعتبار"]
    .sum()
    .reset_index()
)

dastagah_totals = (
    by_company.groupby("دستگاه اجرایی مرتبط")["اعتبار"]
    .sum()
    .sort_values(ascending=False)
)
dastagah_totals = dastagah_totals[dastagah_totals > 0]
dastagah_order = dastagah_totals.index.tolist()

rows = []
for dastagah in dastagah_order:
    group = (
        by_company[by_company["دستگاه اجرایی مرتبط"] == dastagah]
        .sort_values("اعتبار", ascending=False)
        .reset_index(drop=True)
    )
    total = dastagah_totals.loc[dastagah]
    top_companies = group.head(4)

    for rank, row in enumerate(top_companies.itertuples(index=False), start=1):
        rows.append({
            "دستگاه اجرایی مرتبط": dastagah,
            "نام مشمول": row[1],
            "رتبه": rank,
            "اعتبار": row[2],
            "سهم": row[2] / total * 100,
        })

    other_total = group.iloc[4:]["اعتبار"].sum()
    rows.append({
        "دستگاه اجرایی مرتبط": dastagah,
        "نام مشمول": "سایرین",
        "رتبه": 5,
        "اعتبار": other_total,
        "سهم": other_total / total * 100,
    })

table_source = pd.DataFrame(rows)
df_rows = pd.DataFrame({
    "دستگاه": table_source["دستگاه اجرایی مرتبط"],
    "نام_مشمول": table_source["نام مشمول"],
    "اعتبار": table_source["اعتبار"],
    "سهم": table_source["سهم"],
})

# ── خروجی اکسل ───────────────────────────────────────────────────────────────
OUT_FILE = TABLES_DIR / "table09_dastagah_company.xlsx"

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "جزئیات شرکت‌های برتر"
ws.sheet_view.rightToLeft = True

headers = [
    "دستگاه اجرایی مرتبط",
    "نام مشمول",
    "اعتبار تکلیفی (میلیارد ریال)",
    "سهم از دستگاه",
]
header_fill = PatternFill("solid", fgColor="1F4E79")
header_font = Font(bold=True, color="FFFFFF", size=11, name="Vazirmatn")
header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
thin_side = Side(style="thin", color="DDDDDD")
cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

for col_idx, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col_idx, value=h)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = header_align
    cell.border = cell_border

odd_fill = PatternFill("solid", fgColor="FFFFFF")
even_fill = PatternFill("solid", fgColor="F2F2F2")
sayr_fill = PatternFill("solid", fgColor="EDEDED")
center_align = Alignment(horizontal="center", vertical="center")

for r_idx, row in enumerate(df_rows.itertuples(index=False), start=2):
    is_sayr = row.نام_مشمول == "سایرین"
    fill = sayr_fill if is_sayr else (odd_fill if r_idx % 2 else even_fill)
    values = [row.دستگاه, row.نام_مشمول, round(row.اعتبار, 1), round(row.سهم, 1)]
    for c_idx, val in enumerate(values, 1):
        cell = ws.cell(row=r_idx, column=c_idx, value=val)
        cell.fill = fill
        cell.font = Font(size=10, name="Vazirmatn", italic=is_sayr)
        cell.alignment = center_align
        cell.border = cell_border

col_widths = [30, 40, 28, 16]
for i, width in enumerate(col_widths, 1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width
ws.row_dimensions[1].height = 30

for row in ws.iter_rows(min_row=2, min_col=3, max_col=4):
    for cell in row:
        cell.number_format = '#,##0.0'

wb.save(OUT_FILE)
print(f"Table 09 saved → {OUT_FILE}")
