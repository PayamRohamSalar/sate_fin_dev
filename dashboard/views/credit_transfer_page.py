"""Render the credit-transfer (انتقال اعتبار) report/edit/add page.

One page covering all three needs for the ``credit_transfers`` table:
add a new transfer, view the existing ones as a filterable report (with an
Excel export), and edit a selected one -- so the workflow doesn't need to
jump between separate pages for viewing and editing.
"""

import io

import pandas as pd
import requests
import streamlit as st

API_URL = "http://127.0.0.1:8000/api/v1"

ADD_FORM_KEY = "transfer_add_form"
ADD_FORM_WIDGET_KEYS = (
    "transfer_add_source_employer",
    "transfer_add_destination_employer",
    "transfer_add_search_query",
    "transfer_add_contract_choice",
    "transfer_add_fiscal_year",
    "transfer_add_amount",
    "transfer_add_broker_name",
    "transfer_add_documentation",
)

DEFAULT_DESTINATION_EMPLOYER_CODE = "128342"

MILLION = 1_000_000

# Same column structure as the funding-source report finalized with Payam
# (scripts/report_institute_contract_funding.py, "تفصیلی" sheet), all money
# columns in میلیون ریال -- plus the contract code and the few record fields
# that are useful when working from this form (شناسه for editing, سال مالی,
# کارگزار, توضیحات).
COL_ID = "شناسه"
COL_CONTRACT_CODE = "کد قرارداد"
COL_CONTRACT_TITLE = "عنوان قرارداد"
COL_CONTRACT_AMOUNT = "مبلغ قرارداد (میلیون ریال)"
COL_SOURCE = "کارفرمای مبدأ"
COL_SOURCE_CODE = "کد کارفرمای مبدأ"
COL_TRANSFER = "مبلغ انتقالی بابت این قرارداد (میلیون ریال)"
COL_BEFORE = "مانده تجمیعی کارفرمای مبدأ، بدون انتقال اعتبار (میلیون ریال)"
COL_AFTER = "مانده کارفرمای مبدأ پس از همهٔ انتقالی‌هایش (میلیون ریال)"
COL_DESTINATION = "کارفرمای مقصد"
COL_FISCAL_YEAR = "سال مالی"
COL_BROKER = "کارگزار"
COL_DOCUMENTATION = "توضیحات"

REPORT_TABLE_COLUMNS = [
    COL_CONTRACT_CODE,
    COL_CONTRACT_TITLE,
    COL_CONTRACT_AMOUNT,
    COL_SOURCE,
    COL_SOURCE_CODE,
    COL_TRANSFER,
    COL_BEFORE,
    COL_AFTER,
    COL_DESTINATION,
    COL_FISCAL_YEAR,
    COL_BROKER,
    COL_DOCUMENTATION,
    COL_ID,
]
MONEY_COLUMNS = [COL_CONTRACT_AMOUNT, COL_TRANSFER, COL_BEFORE, COL_AFTER]

SUMMARY_COLUMNS = [
    COL_CONTRACT_CODE,
    COL_CONTRACT_TITLE,
    COL_CONTRACT_AMOUNT,
    "تعداد کارفرمای تأمین‌کننده",
    "مجموع انتقالی دریافتی (میلیون ریال)",
    "کسری تأمین (میلیون ریال)",
    "وضعیت تأمین",
]

SELECTED_TRANSFER_ID_KEY = "transfer_edit_selected_id"
EDIT_LOAD_SEQ_KEY = "transfer_edit_load_seq"

st.markdown(
    """
    <style>
    @import url('https://fonts.googleapis.com/css2?family=Vazirmatn:wght@100..900&display=swap');
    * { font-family: 'Vazirmatn', sans-serif !important; }
    .block-container, .stTextInput, .stNumberInput, .stSelectbox,
    .stTextArea, p, div {
        direction: rtl !important;
        text-align: right !important;
    }
    [data-testid="stSidebarNav"] span {
        direction: rtl !important;
        text-align: right !important;
        font-size: 1.05rem !important;
    }
    div[data-testid="InputInstructions"] { display: none !important; }
    h1 { margin-top: 0.2rem !important; margin-bottom: 0.2rem !important; padding: 0 !important; }
    h2, h3 { margin-top: 0.15rem !important; margin-bottom: 0.15rem !important; padding: 0 !important; }
    hr { margin-top: 0.25rem !important; margin-bottom: 0.25rem !important; }
    table, th, td, .stTable, [data-testid="stTable"] * {
        font-family: 'Vazirmatn', sans-serif !important;
        font-size: 15px !important;
        text-align: center !important;
        direction: rtl !important;
    }
    th { background-color: #f0f2f6 !important; border-bottom: 2px solid #4a4a4a !important; color: black !important; }
    </style>
    """,
    unsafe_allow_html=True,
)

st.title("🔄 مدیریت انتقال اعتبار بین کارفرمایان")
st.caption(
    "تخصیص مانده اعتبار مصرف‌نشده یک کارفرما به یک قرارداد مشخص (معمولاً از "
    "طریق دبیرخانه)، بدون تغییر مستقیم جداول واریزی/پرداخت."
)


def _response_detail(response: requests.Response) -> str:
    """Extract a user-facing API error detail from JSON or text responses."""
    try:
        response_data = response.json()
    except ValueError:
        response_data = None

    if isinstance(response_data, dict):
        detail = response_data.get("detail")
        if isinstance(detail, str) and detail.strip():
            return detail
        if isinstance(detail, list):
            messages = [
                item.get("msg", str(item)) if isinstance(item, dict) else str(item)
                for item in detail
            ]
            return " ".join(messages)
    elif isinstance(response_data, str) and response_data.strip():
        return response_data

    return f"خطا در پردازش درخواست (کد وضعیت: {response.status_code})."


def _handle_request_error(exc: Exception, context: str) -> None:
    """Render a localized error for the common request-failure cases."""
    if isinstance(exc, requests.exceptions.ConnectionError):
        st.error(
            "خطا: امکان اتصال به سرور بک‌اند وجود ندارد. لطفاً مطمئن شوید "
            "سرور اصلی (FastAPI) روی پورت 8000 در حال اجرا است."
        )
    elif isinstance(exc, requests.exceptions.Timeout):
        st.error("خطا: پاسخ از سرور بک‌اند بیش از حد طول کشید. لطفاً مجدداً تلاش کنید.")
    else:
        st.error(f"خطای پیش‌بینی‌نشده {context}: {exc}")


def _safe_lookup(path: str) -> list[dict]:
    """Fetch a lookup list and display a localized error when unavailable."""
    try:
        response = requests.get(f"{API_URL}{path}", timeout=5)
        if response.status_code == 200:
            return response.json()
        st.error(f"خطا در دریافت اطلاعات ({path}): {response.status_code}")
    except Exception as exc:
        _handle_request_error(exc, "در دریافت اطلاعات")
    return []


def _format_option(item: dict, value_key: str) -> str:
    """Format lookup entries as a human-readable code and name pair."""
    return f"{item[value_key]} - {item['name']}"


def _format_rials(amount) -> str:
    """Format a numeric amount with thousand separators and a Rial suffix."""
    return f"{amount:,.0f} ریال"


def _optional_text(value: str) -> str | None:
    """Return stripped optional form text or ``None`` when empty."""
    normalized_value = value.strip()
    return normalized_value or None


def _index_of(options: list[str], current_value: str | None) -> int:
    """Return a dropdown's index for a stored FK value, defaulting to 0."""
    if current_value in options:
        return options.index(current_value)
    return 0


def _fetch_combined_balance(employer_code: str) -> float | None:
    """Combined 1402+1404 balance for an employer -- what a credit transfer
    is actually checked against (see ``calculate_combined_pool_balance`` in
    ``services/financial_engine.py``: transfers no longer draw from a single
    fiscal year's pool, since a data error in one year's figures could make
    a perfectly legitimate transfer look like it exceeds that year alone).
    Fetched as two single-year calls against the existing balance endpoint
    and summed client-side, rather than adding a new backend route.
    """
    total = 0.0
    for year in (1402, 1404):
        try:
            response = requests.get(
                f"{API_URL}/employers/{employer_code}/balance",
                params={"funding_year": year},
                timeout=5,
            )
            if response.status_code != 200:
                st.error(f"❌ خطا در دریافت موجودی کارفرما: {_response_detail(response)}")
                return None
            total += response.json()["total_balance"]
        except Exception as exc:
            _handle_request_error(exc, "در دریافت موجودی کارفرما")
            return None
    return total


def _fetch_transfers(*, fiscal_year: int | None, employer_code: str | None) -> list[dict] | None:
    """Fetch the filtered credit-transfer report, handling errors gracefully."""
    params = {}
    if fiscal_year:
        params["fiscal_year"] = fiscal_year
    if employer_code:
        params["employer_code"] = employer_code
    try:
        response = requests.get(f"{API_URL}/credit-transfers/", params=params, timeout=10)
        if response.status_code == 200:
            return response.json()
        st.error(f"❌ خطا در دریافت گزارش: {_response_detail(response)}")
    except Exception as exc:
        _handle_request_error(exc, "در دریافت گزارش انتقال اعتبار")
    return None


def _to_million(value) -> int | None:
    """Rials -> whole میلیون ریال (``None`` stays ``None``)."""
    if value is None:
        return None
    return round(float(value) / MILLION)


def _to_report_row(item: dict) -> dict:
    """Map one enriched transfer record onto the report's Persian headers."""
    return {
        COL_CONTRACT_CODE: item["contract_tracking_code"],
        COL_CONTRACT_TITLE: item.get("contract_title") or "-",
        COL_CONTRACT_AMOUNT: _to_million(item.get("contract_amount")),
        COL_SOURCE: item.get("source_employer_name") or item["source_employer_code"],
        COL_SOURCE_CODE: item["source_employer_code"],
        COL_TRANSFER: _to_million(item["amount"]),
        COL_BEFORE: _to_million(item.get("source_balance_before_transfers")),
        COL_AFTER: _to_million(item.get("source_balance_after_transfers")),
        COL_DESTINATION: item.get("destination_employer_name") or item["destination_employer_code"],
        COL_FISCAL_YEAR: item["fiscal_year"],
        COL_BROKER: item.get("broker_name") or "-",
        COL_DOCUMENTATION: item.get("documentation") or "-",
        COL_ID: item["id"],
    }


def _build_report_dataframe(items: list[dict]) -> pd.DataFrame:
    # Items arrive newest-first; the report reads better grouped by contract.
    rows = [_to_report_row(item) for item in items]
    df = pd.DataFrame(rows, columns=REPORT_TABLE_COLUMNS)
    return df.sort_values([COL_CONTRACT_CODE, COL_ID], kind="stable").reset_index(drop=True)


def _build_summary_dataframe(items: list[dict]) -> pd.DataFrame:
    """One row per contract: how much of its amount the transfers cover.

    The funded total comes from the API (all active transfers of the
    contract), so it stays correct even when the filters hide some rows.
    """
    per_contract: dict[str, dict] = {}
    for item in items:
        entry = per_contract.setdefault(
            item["contract_tracking_code"],
            {"item": item, "sources": set()},
        )
        entry["sources"].add(item["source_employer_code"])

    rows = []
    for code in sorted(per_contract):
        item = per_contract[code]["item"]
        amount = item.get("contract_amount")
        funded = item.get("contract_funded_total") or 0
        shortfall = None if amount is None else amount - funded
        if shortfall is None:
            status = "-"
        elif abs(shortfall) < 1:
            status = "کامل"
        elif shortfall < 0:
            status = "بیش از نیاز"
        else:
            status = "ناقص"
        rows.append({
            COL_CONTRACT_CODE: code,
            COL_CONTRACT_TITLE: item.get("contract_title") or "-",
            COL_CONTRACT_AMOUNT: _to_million(amount),
            "تعداد کارفرمای تأمین‌کننده": len(per_contract[code]["sources"]),
            "مجموع انتقالی دریافتی (میلیون ریال)": _to_million(funded),
            "کسری تأمین (میلیون ریال)": _to_million(shortfall),
            "وضعیت تأمین": status,
        })
    return pd.DataFrame(rows, columns=SUMMARY_COLUMNS)


def _style_excel_sheet(ws, df: pd.DataFrame) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    ws.sheet_view.rightToLeft = True
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    for col_idx, col_name in enumerate(df.columns, start=1):
        header = ws.cell(row=1, column=col_idx)
        header.fill = header_fill
        header.font = Font(bold=True, color="FFFFFF")
        header.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        is_money = col_name in MONEY_COLUMNS or col_name.endswith("(میلیون ریال)")
        for row_idx in range(2, len(df) + 2):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if is_money:
                cell.number_format = "#,##0"
        longest = max([len(str(col_name))] + [len(str(v)) for v in df[col_name]])
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(longest + 3, 12), 55)
    ws.row_dimensions[1].height = 45
    ws.freeze_panes = "A2"


def _to_excel_bytes(detail_df: pd.DataFrame, summary_df: pd.DataFrame) -> bytes:
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        for df, sheet_name in (
            (detail_df, "تفصیلی (قرارداد × کارفرما)"),
            (summary_df, "خلاصهٔ تأمین"),
        ):
            df.to_excel(writer, index=False, sheet_name=sheet_name)
            _style_excel_sheet(writer.sheets[sheet_name], df)
    output.seek(0)
    return output.getvalue()


employers = _safe_lookup("/employers/lookup")
if not employers:
    st.warning(
        "برای ادامه، ابتدا باید حداقل یک کارفرما (غیر زیرمجموعه) در سیستم موجود باشد."
    )
    st.stop()

employer_options = {
    item["classification_code"]: _format_option(item, "classification_code")
    for item in employers
}
employer_keys = list(employer_options)


def _apply_pending_add_reset() -> None:
    if not st.session_state.pop("transfer_add_reset_pending", False):
        return
    for widget_key in ADD_FORM_WIDGET_KEYS:
        st.session_state.pop(widget_key, None)


def _schedule_add_reset() -> None:
    st.session_state.transfer_add_reset_pending = True
    st.rerun()


_apply_pending_add_reset()

add_success_message = st.session_state.pop("transfer_add_success_message", None)
if add_success_message:
    st.success(add_success_message)
    st.toast("✅ انتقال اعتبار جدید ثبت شد.", icon="🔄")

edit_success_message = st.session_state.pop("transfer_edit_success_message", None)
if edit_success_message:
    st.success(edit_success_message)
    st.toast("✅ انتقال اعتبار ویرایش شد.", icon="✏️")

# ─── 1. Add a new transfer ────────────────────────────────────────────────
st.subheader("۱. ثبت انتقال اعتبار جدید")

col_e1, col_e2 = st.columns(2)
with col_e1:
    source_choice = st.selectbox(
        "🏢 کارفرمای مبدأ:",
        options=employer_keys,
        format_func=employer_options.get,
        key="transfer_add_source_employer",
    )
    _source_balance = _fetch_combined_balance(source_choice)
    if _source_balance is not None:
        st.caption(
            f"💰 مانده تجمیعی مبدأ (۱۴۰۲ + ۱۴۰۴ با هم): {_format_rials(_source_balance)} ریال"
        )
with col_e2:
    destination_default_index = _index_of(employer_keys, DEFAULT_DESTINATION_EMPLOYER_CODE)
    destination_choice = st.selectbox(
        "🏛️ کارفرمای مقصد:",
        options=employer_keys,
        format_func=employer_options.get,
        index=destination_default_index,
        key="transfer_add_destination_employer",
    )

search_query = st.text_input(
    "🔍 جستجوی قرارداد مقصد انتقال (کد رهگیری یا بخشی از عنوان):",
    key="transfer_add_search_query",
)
selected_contract = None
if search_query:
    try:
        response = requests.get(
            f"{API_URL}/contracts/search", params={"q": search_query}, timeout=5
        )
        if response.status_code == 200 and response.json():
            contracts_found = response.json()
            options = {
                f"{c['contract_tracking_code']} - {c['title'][:40]}": c
                for c in contracts_found
            }
            choice = st.selectbox(
                "قرارداد مورد نظر را انتخاب کنید:",
                list(options.keys()),
                key="transfer_add_contract_choice",
            )
            selected_contract = options[choice]
        else:
            st.warning("هیچ قراردادی با این مشخصات یافت نشد.")
    except Exception as exc:
        _handle_request_error(exc, "در جستجوی قرارداد")

col_r1, col_r2 = st.columns(2)
with col_r1:
    add_fiscal_year = st.number_input(
        "📅 سال مالی مبدأ اعتبار (شمسی):",
        min_value=1398,
        max_value=1410,
        value=1405,
        step=1,
        key="transfer_add_fiscal_year",
    )
with col_r2:
    add_amount = st.number_input(
        "💵 مبلغ انتقالی (ریال):",
        min_value=0,
        step=1_000_000,
        format="%d",
        value=0,
        key="transfer_add_amount",
    )
col_b1, col_b2 = st.columns(2)
with col_b1:
    add_broker_name = st.text_input("کارگزار (اختیاری):", key="transfer_add_broker_name")
with col_b2:
    add_documentation = st.text_input(
        "توضیحات/مستندات (اختیاری):", key="transfer_add_documentation"
    )

if st.button("➕ ثبت انتقال اعتبار", type="primary"):
    if selected_contract is None:
        st.error("لطفاً ابتدا یک قرارداد مقصد را جستجو و انتخاب کنید.")
    elif add_amount <= 0:
        st.error("مبلغ انتقالی باید بیشتر از صفر باشد.")
    else:
        payload = {
            "source_employer_code": source_choice,
            "destination_employer_code": destination_choice,
            "contract_tracking_code": selected_contract["contract_tracking_code"],
            "fiscal_year": int(add_fiscal_year),
            "amount": int(add_amount),
            "broker_name": _optional_text(add_broker_name),
            "documentation": _optional_text(add_documentation),
        }
        try:
            response = requests.post(
                f"{API_URL}/credit-transfers/", json=payload, timeout=10
            )
            if response.status_code == 201:
                data = response.json()
                st.session_state.transfer_add_success_message = (
                    "🎉 انتقال اعتبار با موفقیت ثبت شد. "
                    f"از {data.get('source_employer_name')} به "
                    f"{data.get('destination_employer_name')} | "
                    f"مبلغ: {_format_rials(data.get('amount', 0))}"
                )
                _schedule_add_reset()
            else:
                st.error(f"❌ خطا در ثبت انتقال اعتبار: {_response_detail(response)}")
        except Exception as exc:
            _handle_request_error(exc, "در ثبت انتقال اعتبار")

st.markdown("---")

# ─── 2. Report: filter + view + Excel export ─────────────────────────────
st.subheader("۲. گزارش انتقال‌های ثبت‌شده")

col_f1, col_f2 = st.columns(2)
with col_f1:
    filter_fiscal_year = st.selectbox(
        "فیلتر بر اساس سال مالی:",
        options=[None, 1402, 1403, 1404, 1405, 1406],
        format_func=lambda y: "همه سال‌ها" if y is None else str(y),
        key="transfer_report_fiscal_year",
    )
with col_f2:
    filter_employer_options = {None: "همه کارفرمایان"}
    filter_employer_options.update(employer_options)
    filter_employer_choice = st.selectbox(
        "فیلتر بر اساس کارفرما (مبدأ یا مقصد):",
        options=list(filter_employer_options),
        format_func=filter_employer_options.get,
        key="transfer_report_employer",
    )

transfers = _fetch_transfers(
    fiscal_year=filter_fiscal_year, employer_code=filter_employer_choice
)

if transfers:
    report_df = _build_report_dataframe(transfers)
    summary_df = _build_summary_dataframe(transfers)

    def _format_million(value) -> str:
        return "-" if pd.isna(value) else f"{value:,.0f}"

    display_df = report_df.copy()
    for money_col in MONEY_COLUMNS:
        display_df[money_col] = display_df[money_col].apply(_format_million)
    st.caption(
        "مبالغ به میلیون ریال. «مانده بدون انتقال» = واریزی منهای هزینه‌کرد (۱۴۰۲+۱۴۰۴)؛ "
        "«مانده پس از انتقال‌ها» شامل تمام انتقال‌های آن کارفرما (نه فقط همین قرارداد) است."
    )
    st.table(display_df)

    st.subheader("خلاصهٔ تأمین هر قرارداد")
    summary_display = summary_df.copy()
    for money_col in (
        COL_CONTRACT_AMOUNT,
        "مجموع انتقالی دریافتی (میلیون ریال)",
        "کسری تأمین (میلیون ریال)",
    ):
        summary_display[money_col] = summary_display[money_col].apply(_format_million)
    st.table(summary_display)

    st.download_button(
        label="📥 دانلود گزارش به صورت اکسل",
        data=_to_excel_bytes(report_df, summary_df),
        file_name="credit_transfers_report.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

    st.markdown("---")

    # ─── 3. Edit a selected transfer ──────────────────────────────────────
    st.subheader("۳. ویرایش یک ردیف")

    edit_options = {
        f"#{t['id']} — {t.get('source_employer_name') or t['source_employer_code']} → "
        f"{t.get('destination_employer_name') or t['destination_employer_code']} "
        f"({_format_rials(t['amount'])}, {t['fiscal_year']})": t
        for t in transfers
    }
    edit_choice = st.selectbox(
        "ردیف مورد نظر برای ویرایش را انتخاب کنید:",
        list(edit_options),
        key="transfer_edit_choice",
    )
    selected_transfer = edit_options[edit_choice]
    transfer_id = selected_transfer["id"]

    if st.session_state.get(SELECTED_TRANSFER_ID_KEY) != transfer_id:
        st.session_state[SELECTED_TRANSFER_ID_KEY] = transfer_id
        # Bump the load sequence so widget keys are unique per fetch, not
        # just per transfer_id -- otherwise re-selecting the same row after
        # editing it once would reuse stale cached widget values instead of
        # the freshly fetched ones (same fix as contract_edit_page.py).
        st.session_state[EDIT_LOAD_SEQ_KEY] = (
            st.session_state.get(EDIT_LOAD_SEQ_KEY, 0) + 1
        )

    form_instance_key = f"{transfer_id}_{st.session_state.get(EDIT_LOAD_SEQ_KEY, 0)}"

    with st.form(f"transfer_edit_form_{form_instance_key}", clear_on_submit=False):
        col_ee1, col_ee2 = st.columns(2)
        with col_ee1:
            edit_source_choice = st.selectbox(
                "🏢 کارفرمای مبدأ:",
                options=employer_keys,
                format_func=employer_options.get,
                index=_index_of(employer_keys, selected_transfer["source_employer_code"]),
                key=f"edit_source_{form_instance_key}",
            )
        with col_ee2:
            edit_destination_choice = st.selectbox(
                "🏛️ کارفرمای مقصد:",
                options=employer_keys,
                format_func=employer_options.get,
                index=_index_of(
                    employer_keys, selected_transfer["destination_employer_code"]
                ),
                key=f"edit_destination_{form_instance_key}",
            )

        edit_contract_code = st.text_input(
            "کد قرارداد مقصد:",
            value=selected_transfer["contract_tracking_code"],
            key=f"edit_contract_{form_instance_key}",
        )
        if selected_transfer.get("contract_title"):
            st.caption(f"عنوان فعلی قرارداد: {selected_transfer['contract_title']}")

        col_er1, col_er2 = st.columns(2)
        with col_er1:
            edit_fiscal_year = st.number_input(
                "📅 سال مالی مبدأ اعتبار (شمسی):",
                min_value=1398,
                max_value=1410,
                value=int(selected_transfer["fiscal_year"]),
                step=1,
                key=f"edit_fiscal_year_{form_instance_key}",
            )
        with col_er2:
            edit_amount = st.number_input(
                "💵 مبلغ انتقالی (ریال):",
                min_value=0,
                step=1_000_000,
                format="%d",
                value=int(selected_transfer["amount"]),
                key=f"edit_amount_{form_instance_key}",
            )
        col_eb1, col_eb2 = st.columns(2)
        with col_eb1:
            edit_broker_name = st.text_input(
                "کارگزار (اختیاری):",
                value=selected_transfer.get("broker_name") or "",
                key=f"edit_broker_{form_instance_key}",
            )
        with col_eb2:
            edit_documentation = st.text_input(
                "توضیحات/مستندات (اختیاری):",
                value=selected_transfer.get("documentation") or "",
                key=f"edit_documentation_{form_instance_key}",
            )

        edit_submitted = st.form_submit_button("💾 ذخیره تغییرات", type="primary")

    if edit_submitted:
        if not edit_contract_code.strip():
            st.error("لطفاً کد قرارداد را وارد کنید.")
        elif edit_amount <= 0:
            st.error("مبلغ انتقالی باید بیشتر از صفر باشد.")
        else:
            edit_payload = {
                "source_employer_code": edit_source_choice,
                "destination_employer_code": edit_destination_choice,
                "contract_tracking_code": edit_contract_code.strip(),
                "fiscal_year": int(edit_fiscal_year),
                "amount": int(edit_amount),
                "broker_name": _optional_text(edit_broker_name),
                "documentation": _optional_text(edit_documentation),
            }
            try:
                response = requests.put(
                    f"{API_URL}/credit-transfers/{transfer_id}",
                    json=edit_payload,
                    timeout=10,
                )
                if response.status_code == 200:
                    data = response.json()
                    st.session_state.transfer_edit_success_message = (
                        f"🎉 انتقال اعتبار #{data.get('id')} با موفقیت ویرایش شد. "
                        f"از {data.get('source_employer_name')} به "
                        f"{data.get('destination_employer_name')} | "
                        f"مبلغ: {_format_rials(data.get('amount', 0))}"
                    )
                    st.session_state.pop(SELECTED_TRANSFER_ID_KEY, None)
                    st.rerun()
                else:
                    st.error(f"❌ خطا در ذخیره تغییرات: {_response_detail(response)}")
            except Exception as exc:
                _handle_request_error(exc, "در ذخیره تغییرات")
elif transfers is not None:
    st.info("هیچ انتقال اعتباری با این فیلترها یافت نشد.")
